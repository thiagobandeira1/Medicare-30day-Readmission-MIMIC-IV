# -*- coding: utf-8 -*-
"""Value-level consistency report for the v6 manuscript vs canonical JSON."""
import json, re
from pathlib import Path
from docx import Document

PUB = Path(__file__).resolve().parent.parent
O6 = PUB / "medicare-30day-readmission-mimic-iv" / "results_v6"
O5 = PUB / "medicare-30day-readmission-mimic-iv" / "results_v5"
FM = json.loads((O6 / "final_model_v6.json").read_text())
EX = FM["v6_extras"]
CMP = json.loads((O5 / "model_comparison.json").read_text())
BD = json.loads((O5 / "boundary_v5.json").read_text())

doc = Document(str(PUB / "Paper JMIR AI Submission v6.docx"))
text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            text += "\n" + c.text

report = []


def check(name, value, expect_absent=False):
    n = text.count(value)
    ok = (n == 0) if expect_absent else (n > 0)
    report.append(f"{'OK  ' if ok else 'FAIL'} {name}: {value!r} x{n}")
    return ok


VH = FM["validation_hierarchy"]
fx = EX["fixed31_vs_fixed142_paired_cv"]
hd = EX["highdraw_contrasts_5000"]
tie = EX["tie_rule_sensitivity"]
MP = FM["metric_panel"]
ok = True
# canonical values must appear
for name, v in [
    ("primary OOF", f"{VH['primary_oof_procedure_dev_only']['auroc']:.4f}"),
    ("secondary CV", f"{VH['secondary_consensus_cv_dev_only']['mean']:.4f}"),
    ("tertiary test", f"{MP['auroc']['point']:.4f}"),
    ("n features", f"{FM['n_features']}-feature"),
    ("threshold", f"{FM['threshold_validation']:.3f}"),
    ("slope", f"{MP['slope']['point']:.2f}"),
    ("ECE", f"{MP['ece']['point']:.3f}"),
    ("LACE", f"{FM['baselines']['lace_auroc']:.4f}"),
    ("HOSPITAL12", f"{FM['baselines']['hospital12_auroc']:.4f}"),
    ("WB gap", f"{FM['fairness']['white_minus_black']['point']:.4f}"),
    ("fixed diff", f"{fx['paired_diff_31_minus_142']:+.4f}"),
    ("fixed31 CV", f"{fx['auroc_fixed31_cv']:.4f}"),
    ("fixed142 CV", f"{fx['auroc_fixed142_cv']:.4f}"),
    ("tie C alt", f"{tie['readmission_first']['harrell_c']:.4f}"),
    ("boundary safe", f"{BD['auroc_safe']:.4f}"),
    ("AFT C", f"{FM['survival']['aft_harrell_c']:.4f}"),
    ("sameday refit", f"{FM['sensitivity']['sameday_as_readmission']['test_auroc_refit']:.4f}"),
    ("unplanned refit", f"{FM['sensitivity']['unplanned_only']['test_auroc_refit']:.4f}"),
    ("nonelective", f"{FM['sensitivity']['nonelective_index_only']['test_auroc']:.4f}"),
    ("72%", "72%"), ("8%", "8%"), ("20%", "20%"),
]:
    ok &= check(name, v)
# stale/forbidden strings must be absent
for name, v in [
    ("80/10/10", "80/10/10"),
    ("6-month proxy as available", "6-month window as available proxy"),
    ("indistinguishable", "indistinguishab"),
    ("prespecified conservative tie rule", "a prespecified conservative rule"),
    ("stale 33-feature", "33-feature"),
    ("stale 143", "143-feature"),
    ("proved provenance", "showed it was built by"),
    ("validation-derived threshold", "validation-derived threshold"),
    ("paired WB claim", "identical observations - out-of-fold comparators and the White-Black"),
    ("consensus-31 mislabel of OOF", f"consensus-{FM['n_features']}) "
     f"{VH['primary_oof_procedure_dev_only']['auroc']:.4f}"),
]:
    ok &= check(name, v, expect_absent=True)
# every P value in text must be <.001 or =.xxx form (no bare P<.0001 etc)
badp = re.findall(r"P[<=]\.\d{4,}", text)
report.append(f"{'OK  ' if not badp else 'FAIL'} no over-precise P values: {badp}")
ok &= not badp
figs = sorted(set(re.findall(r"Figure (\d)", text)))
tabs = sorted(set(re.findall(r"Table (\d)", text)))
report.append(f"figures referenced: {figs}; tables referenced: {tabs}")
out = PUB / "CONSISTENCY_REPORT_V6.md"
out.write_text("# v6 value-level consistency report\n\n"
               + "\n".join("- " + r for r in report)
               + f"\n\nOVERALL: {'PASS' if ok else 'FAIL'}\n",
               encoding="utf-8")
print("\n".join(report))
print("OVERALL:", "PASS" if ok else "FAIL")
