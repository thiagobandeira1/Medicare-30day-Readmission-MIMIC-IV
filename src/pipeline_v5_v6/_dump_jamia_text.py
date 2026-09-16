# -*- coding: utf-8 -*-
"""Dump the built JAMIA manuscript and supplement to plain text for review."""
from pathlib import Path
from docx import Document

PUB = Path(__file__).resolve().parent.parent
for src, dst in (("Paper JAMIA Submission FINAL.docx", "_jamia_current.txt"),
                 ("Paper JAMIA Supplementary Material.docx",
                  "_jamia_supp_current.txt")):
    doc = Document(str(PUB / src))
    out = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = p.style.name
        if style.startswith("Heading"):
            out.append(f"\n{'#' * int(style[-1])} {txt}")
        else:
            out.append(txt)
    for i, t in enumerate(doc.tables, 1):
        out.append(f"\n[TABLE {i}]")
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    d = PUB / "presentation-update" / dst
    d.write_text("\n".join(out), encoding="utf-8")
    print(f"wrote {d} ({d.stat().st_size:,} bytes)")
